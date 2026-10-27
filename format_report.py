
from openpyxl import load_workbook
from datetime import date, datetime
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt


REQUIRED_SHEETS = {
    "Fase 1.0 | Loopbaanfasen",
    "Fase 1.1 | Big Five Dimensies",
    "Fase 2.0 | Loopbaanankers",
    "Fase 2.1 | Carriere Clusters",
    "Fase 2.2 | Cultuur analyse",
    "Fase 2.3 | J.C.M.",
    "Gegevens",
    "Rapport",
}

BIG_FIVE_SCORE_FORMULAS = {
    "Extraversie": (
        20,
        (
            ("C4", 1), ("C9", -1), ("C14", 1), ("C19", -1),
            ("C24", 1), ("C29", -1), ("C34", 1), ("C39", -1),
            ("C44", 1), ("C49", -1),
        ),
    ),
    "Altruïsme": (
        14,
        (
            ("D5", -1), ("D10", 1), ("D15", -1), ("D20", 1),
            ("D25", -1), ("D30", 1), ("D35", -1), ("D40", 1),
            ("D45", 1), ("D50", 1),
        ),
    ),
    "Consciëntieusheid": (
        14,
        (
            ("E6", 1), ("E11", -1), ("E16", 1), ("E21", -1),
            ("E26", 1), ("E31", -1), ("E36", 1), ("E41", -1),
            ("E46", 1), ("E51", 1),
        ),
    ),
    "Neuroticisme": (
        38,
        (
            ("F7", -1), ("F12", 1), ("F17", -1), ("F22", 1),
            ("F27", -1), ("F32", -1), ("F37", -1), ("F42", -1),
            ("F47", -1), ("F52", -1),
        ),
    ),
    "Openheid": (
        8,
        (
            ("G8", 1), ("G13", -1), ("G18", 1), ("G23", -1),
            ("G28", 1), ("G33", -1), ("G38", 1), ("G43", 1),
            ("G48", 1), ("G53", 1),
        ),
    ),
}


def _number(value):
    if value is None or value == "":
        return 0

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value

    raise ValueError(
        f"Expected a numeric assessment value, got {value!r}"
    )


def _get_big_five_scores(sheet):
    return {
        trait: base + sum(
            _number(sheet[cell].value) * factor
            for cell, factor in terms
        )
        for trait, (base, terms) in BIG_FIVE_SCORE_FORMULAS.items()
    }


def _get_big_five_interpretations(scores, data_sheet):
    description_rows = {
        "Extraversie": 25,
        "Altruïsme": 26,
        "Consciëntieusheid": 27,
        "Neuroticisme": 28,
        "Openheid": 29,
    }

    interpretations = {}

    for trait, score in scores.items():
        description_column = (
            "C" if score <= 17
            else "B" if score <= 26
            else "A"
        )

        description = data_sheet[
            f"{description_column}{description_rows[trait]}"
        ].value

        if not description:
            raise ValueError(
                f"Geen score-interpretatie gevonden voor {trait}."
            )

        interpretations[trait] = description

    return interpretations


def _get_named_descriptions(data_sheet, first_row, last_row):
    descriptions = {}

    for row in data_sheet.iter_rows(
        min_row=first_row,
        max_row=last_row,
        min_col=1,
        max_col=1,
        values_only=True,
    ):
        value = row[0]

        if value and "|" in str(value):
            name, description = str(value).split("|", 1)
            descriptions[name.strip()] = description.strip()

    return descriptions


def _get_culture_descriptions(data_sheet):
    descriptions = {}

    for row in data_sheet.iter_rows(
        min_row=93,
        max_row=96,
        min_col=1,
        max_col=1,
        values_only=True,
    ):
        value = row[0]

        if not value:
            continue

        lines = str(value).strip().splitlines()
        name = lines[0].strip()

        if name.startswith("De "):
            name = name[3:]

        description = "\n".join(
            line.strip()
            for line in lines[1:]
            if line.strip()
        )

        if name and description:
            descriptions[name] = description

    return descriptions


def _get_career_phase_descriptions(sheet):
    return {
        str(sheet.cell(row=row, column=2).value).strip():
        str(sheet.cell(row=row, column=3).value).strip()
        for row in range(3, sheet.max_row + 1)
        if (
            sheet.cell(row=row, column=2).value
            and sheet.cell(row=row, column=3).value
        )
    }


def _with_description(name, descriptions, section, aliases=()):
    for candidate in (name, *aliases):
        description = descriptions.get(candidate)

        if description:
            return f"{name}\n{description}"

    raise ValueError(
        f"Geen omschrijving gevonden voor {name!r} in {section}."
    )


def _get_highest_anchors(sheet):
    anchor_names = {
        "V": "Omhoog komen",
        "W": "Veilig voelen",
        "X": "Vrij zijn",
        "Y": "Balans vinden",
        "Z": "Uitdaging zoeken",
    }

    scores = {
        code: sum(
            _number(sheet[f"{column}{row}"].value)
            for row in range(4, 100)
        )
        for code, column in zip(anchor_names, "CDEFG")
    }

    return [
        anchor_names[code]
        for code, score in sorted(
            scores.items(),
            key=lambda item: item[1],
            reverse=True,
        )
        if score > 0
    ][:2]


def _get_highest_clusters(sheet):
    scores = []

    for cluster_id in range(1, 17):
        first_row = 4 + (cluster_id - 1) * 8
        name = sheet[f"I{first_row}"].value

        score = sum(
            _number(sheet[f"C{row}"].value)
            for row in range(first_row, first_row + 7)
        )

        score += sum(
            _number(sheet[f"E{row}"].value)
            for row in range(first_row, first_row + 5)
        )

        score += sum(
            _number(sheet[f"G{row}"].value)
            for row in range(first_row, first_row + 5)
        )

        if name and score > 0:
            scores.append((str(name).strip(), score))

    return [
        name
        for name, _ in sorted(
            scores,
            key=lambda item: item[1],
            reverse=True,
        )[:2]
    ]


def _get_highest_cultures(sheet):
    scores = []

    for first_row, name_row in zip(
        (2, 7, 12, 17),
        (2, 8, 13, 18),
    ):
        name = sheet[f"E{name_row}"].value

        score = sum(
            _number(sheet[f"C{row}"].value)
            for row in range(first_row, first_row + 4)
        )

        if name and score > 0:
            scores.append((str(name).strip(), score))

    return [
        name
        for name, _ in sorted(
            scores,
            key=lambda item: item[1],
            reverse=True,
        )[:2]
    ]


def _get_client_career_phase(client):
    value = client.get("Date of Birth")

    if isinstance(value, date):
        birth_date = value
    else:
        birth_date = None

        for date_format in (
            "%d-%m-%Y",
            "%d-%m-%y",
            "%d/%m/%Y",
            "%Y-%m-%d",
        ):
            try:
                birth_date = datetime.strptime(
                    str(value), date_format
                ).date()
                break
            except ValueError:
                continue

    if birth_date is None:
        raise ValueError(
            "De geboortedatum van de cliënt ontbreekt "
            "of heeft een ongeldig formaat."
        )

    today = date.today()

    age = (
        today.year
        - birth_date.year
        - (
            (today.month, today.day)
            < (birth_date.month, birth_date.day)
        )
    )

    if 15 <= age <= 24:
        return "Verkenning"

    if 25 <= age <= 44:
        return "Keuze"

    if 45 <= age <= 64:
        return "Onderhoud"

    if age >= 65:
        return "Terugtrekking"

    raise ValueError(
        "De cliënt moet minimaal 15 jaar oud zijn "
        "om een loopbaanfase te bepalen."
    )


def _write_after_label(sheet, label, value, occurrence=1):
    matches = [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if (
            cell.value is not None
            and str(cell.value).strip() == label
        )
    ]

    if len(matches) < occurrence:
        raise ValueError(
            f"Rapport-label niet gevonden: "
            f"{label} (voorkomen {occurrence})"
        )

    cell = matches[occurrence - 1]

    sheet.cell(
        row=cell.row,
        column=cell.column + 1,
        value=value,
    )


def generate_report(excel_path, client=None):
    wb = load_workbook(excel_path)

    missing_sheets = REQUIRED_SHEETS.difference(wb.sheetnames)

    if missing_sheets:
        wb.close()
        raise ValueError(
            "Dit is geen resultatenbestand. Ontbrekende tabbladen: "
            + ", ".join(sorted(missing_sheets))
        )

    rapport = wb["Rapport"]
    fase11 = wb["Fase 1.1 | Big Five Dimensies"]
    fase20 = wb["Fase 2.0 | Loopbaanankers"]
    fase21 = wb["Fase 2.1 | Carriere Clusters"]
    fase22 = wb["Fase 2.2 | Cultuur analyse"]
    fase23 = wb["Fase 2.3 | J.C.M."]
    fase10 = wb["Fase 1.0 | Loopbaanfasen"]
    gegevens = wb["Gegevens"]

    big_five_scores = _get_big_five_scores(fase11)

    big_five_interpretations = _get_big_five_interpretations(
        big_five_scores,
        gegevens,
    )

    anchor_descriptions = _get_named_descriptions(
        gegevens, 31, 35
    )

    cluster_descriptions = _get_named_descriptions(
        gegevens, 37, 52
    )

    culture_descriptions = _get_culture_descriptions(gegevens)

    career_phase_descriptions = _get_career_phase_descriptions(
        fase10
    )

    # 1. Loopbaanankers
    top_anchors = _get_highest_anchors(fase20)

    anchor_1 = top_anchors[0] if len(top_anchors) > 0 else ""
    anchor_2 = top_anchors[1] if len(top_anchors) > 1 else ""

    # 2. Carrièreclusters
    top_clusters = _get_highest_clusters(fase21)

    cluster_1 = top_clusters[0] if len(top_clusters) > 0 else ""
    cluster_2 = top_clusters[1] if len(top_clusters) > 1 else ""

    # 3. Cultuuranalyse
    top_cultures = _get_highest_cultures(fase22)

    culture_1 = top_cultures[0] if len(top_cultures) > 0 else ""
    culture_2 = top_cultures[1] if len(top_cultures) > 1 else ""

    # 4. J.C.M.-scores
    def get_jcm_scores():
        jcm_scores = {}

        for row in fase23.iter_rows(
            min_row=1,
            max_row=50,
            values_only=False,
        ):
            if row[0].value:
                cell_value = str(row[0].value).strip()

                if cell_value in [
                    "Taakvaardigheid",
                    "Taakidentiteit",
                    "Taakbetekenis",
                    "Autonomie",
                    "Feedback",
                ]:
                    if len(row) > 3 and row[3].value:
                        jcm_scores[cell_value] = str(row[3].value)
                    else:
                        jcm_scores[cell_value] = ""

        return jcm_scores

    jcm_scores = get_jcm_scores()

    # 5. Loopbaanfase
    if client is not None:
        career_phase = _get_client_career_phase(client)

        phase_value = _with_description(
            career_phase,
            career_phase_descriptions,
            "Fase 1.0",
        )

        _write_after_label(
            rapport,
            "Loopbaan fase *",
            phase_value,
        )

    # 6. Big Five
    for trait, interpretation in big_five_interpretations.items():
        _write_after_label(
            rapport,
            f"{trait} *",
            interpretation,
        )

    # 7. Loopbaanankers
    for occurrence, value in enumerate(
        (anchor_1, anchor_2),
        start=1,
    ):
        described_value = (
            _with_description(
                value,
                anchor_descriptions,
                "Gegevens",
            )
            if value
            else value
        )

        _write_after_label(
            rapport,
            f"Hoogste score {occurrence} *",
            described_value,
        )

    # 8. Carrièreclusters
    for label, value in zip(
        ("Hoogste score 1 *", "Hoogste score 2 *"),
        (cluster_1, cluster_2),
    ):
        aliases = (
            ("Milieu, Landbouw, voeding en natuurlijke grondstoffen",)
            if value == "Landbouw, voeding en natuurlijke grondstoffen"
            else ()
        )

        described_value = (
            _with_description(
                value,
                cluster_descriptions,
                "Gegevens",
                aliases,
            )
            if value
            else value
        )

        _write_after_label(
            rapport,
            label,
            described_value,
            occurrence=2,
        )

    # 9. Cultuuranalyse
    for label, value in zip(
        ("Hoogste score 1 *", "Hoogste score 2 *"),
        (culture_1, culture_2),
    ):
        described_value = (
            _with_description(
                value,
                culture_descriptions,
                "Gegevens",
                (f"De {value}",),
            )
            if value
            else value
        )

        _write_after_label(
            rapport,
            label,
            described_value,
            occurrence=3,
        )

    # 10. J.C.M.
    for component, value in jcm_scores.items():
        _write_after_label(
            rapport,
            component,
            value,
        )

    wb.save(excel_path)
    wb.close()

    return excel_path


def generate_word_report(excel_path, client=None):
    wb = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    try:
        if "Rapport" not in wb.sheetnames:
            raise ValueError(
                "Het tabblad Rapport ontbreekt in het resultatenbestand."
            )

        rapport = wb["Rapport"]

        report_rows = [
            [
                cell.value
                for cell in row
                if cell.value is not None
            ]
            for row in rapport.iter_rows()
        ]

    finally:
        wb.close()

    document = Document()

    # Standaard lettertype en tekstgrootte
    for style in document.styles:
        if style.type in (
            WD_STYLE_TYPE.PARAGRAPH,
            WD_STYLE_TYPE.CHARACTER,
        ):
            style.font.name = "Calibri"

    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    # Hoofdtitel: 14 pt
    title = document.add_heading(
        "Loopbaanrapport",
        level=0,
    )
    title.style.font.name = "Calibri"
    title.style.font.size = Pt(14)

    # Alle kopjes: 12 pt
    for style_name in (
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(12)

    # Naam cliënt
    if client and client.get("name"):
        subtitle = document.add_paragraph(
            client["name"],
            style="Subtitle",
        )

        for run in subtitle.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(11)

    big_five_traits = set(BIG_FIVE_SCORE_FORMULAS)

    for values in report_rows:
        values = [
            str(value).strip()
            for value in values
        ]

        if not values:
            continue

        if values == ["* Selecteer optie in uitvouwmenu"]:
            continue

        # Rijen met één cel: kopje of gewone tekst
        if len(values) == 1:
            text = values[0]

            if text.startswith("Fase "):
                heading = document.add_heading(
                    text,
                    level=1,
                )
                heading.style.font.name = "Calibri"
                heading.style.font.size = Pt(12)

            else:
                paragraph = document.add_paragraph(text)

                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

            continue

        label, value = values[:2]
        label = label.rstrip("*").strip()

        # Big Five: kopje van 12 pt en tekst van 11 pt
        if label in big_five_traits:
            heading = document.add_heading(
                label,
                level=2,
            )
            heading.style.font.name = "Calibri"
            heading.style.font.size = Pt(12)

            paragraph = document.add_paragraph(value)

            for run in paragraph.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(11)

        else:
            # Opsommingstekens voor overige rapportregels
            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            label_run = paragraph.add_run(
                f"{label}: "
            )
            label_run.bold = True
            label_run.font.name = "Calibri"
            label_run.font.size = Pt(11)

            value_run = paragraph.add_run(value)
            value_run.font.name = "Calibri"
            value_run.font.size = Pt(11)

    # Word-bestand opslaan naast het Excel-bestand
    word_path = Path(excel_path).with_name(
        f"{Path(excel_path).stem}_report.docx"
    )

    document.save(word_path)

    return str(word_path)


# ============================================
# DEBUG FUNCTION
# ============================================

def debug_excel_structure(excel_path):
    """Print de structuur van het Excel-bestand."""

    wb = load_workbook(excel_path)

    print("\n=== Fase 2.0 | Loopbaanankers ===")

    fase20 = wb["Fase 2.0 | Loopbaanankers"]

    for row_idx, row in enumerate(
        fase20.iter_rows(
            min_row=1,
            max_row=150,
            values_only=False,
        ),
        1,
    ):
        if row[0].value and "Totaal score" in str(row[0].value):
            print(f"Row {row_idx}: A='{row[0].value}'")

            print(
                f"  D (V/Omhoog): "
                f"{row[3].value if len(row) > 3 else 'EMPTY'}"
            )

            print(
                f"  E (W/Veilig): "
                f"{row[4].value if len(row) > 4 else 'EMPTY'}"
            )

            print(
                f"  F (X/Vrij): "
                f"{row[5].value if len(row) > 5 else 'EMPTY'}"
            )

            print(
                f"  G (Y/Balans): "
                f"{row[6].value if len(row) > 6 else 'EMPTY'}"
            )

            print(
                f"  H (Z/Uitdaging): "
                f"{row[7].value if len(row) > 7 else 'EMPTY'}"
            )

    print("\n=== Fase 2.1 | Carriere Clusters ===")

    fase21 = wb["Fase 2.1 | Carriere Clusters"]

    for row_idx, row in enumerate(
        fase21.iter_rows(
            min_row=10,
            max_row=200,
            values_only=False,
        ),
        1,
    ):
        if row[0].value is not None:
            try:
                cluster_num = int(row[0].value)

                if 1 <= cluster_num <= 16 and len(row) > 8:
                    print(
                        f"Cluster {cluster_num}: "
                        f"Score = {row[8].value if row[8].value else 0}"
                    )

            except (ValueError, TypeError):
                pass

    print("\n=== Fase 2.2 | Cultuur analyse ===")

    fase22 = wb["Fase 2.2 | Cultuur analyse"]

    for row_idx, row in enumerate(
        fase22.iter_rows(
            min_row=5,
            max_row=25,
            values_only=False,
        ),
        1,
    ):
        if row[0].value is not None:
            try:
                culture_num = int(row[0].value)

                if 1 <= culture_num <= 4 and len(row) > 4:
                    print(
                        f"Cultuur {culture_num}: "
                        f"Score = {row[4].value if row[4].value else 0}"
                    )

            except (ValueError, TypeError):
                pass

    print("\n=== Fase 2.3 | J.C.M. ===")

    fase23 = wb["Fase 2.3 | J.C.M."]

    for row_idx, row in enumerate(
        fase23.iter_rows(
            min_row=1,
            max_row=50,
            values_only=False,
        ),
        1,
    ):
        if row[0].value:
            val = str(row[0].value).strip()

            if val in [
                "Taakvaardigheid",
                "Taakidentiteit",
                "Taakbetekenis",
                "Autonomie",
                "Feedback",
            ]:
                print(f"Row {row_idx}: {val}")

                if len(row) > 3:
                    print(
                        f"  Answer (Column D): "
                        f"{row[3].value if row[3].value else 'EMPTY'}"
                    )

                if len(row) > 2:
                    print(
                        f"  Question (Column C): "
                        f"{row[2].value if row[2].value else 'EMPTY'}"
                    )

    wb.close()
