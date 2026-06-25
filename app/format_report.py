# app/format_report.py
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import json
import shutil
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from typing import Dict, Any, Optional
import sys
import os

# Import from utils
from app.utils import get_client, get_client_excel_path, get_client_word_path, save_client


def get_base_path():
    """Get the base path for the application (supports .exe and script mode)"""
    if getattr(sys, 'frozen', False):
        return sys._MEIPASS
    else:
        return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


BASE_PATH = get_base_path()

# Excel template path - try multiple locations
EXCEL_TEMPLATE = None
possible_paths = [
    Path("Loopbaan onderzoek 5.0.xlsx"),
    Path(BASE_PATH) / "Loopbaan onderzoek 5.0.xlsx",
    Path(os.path.dirname(BASE_PATH)) / "Loopbaan onderzoek 5.0.xlsx",
    Path(os.path.join(BASE_PATH, "..", "Loopbaan onderzoek 5.0.xlsx")),
]

for path in possible_paths:
    if path.exists():
        EXCEL_TEMPLATE = path
        break

if EXCEL_TEMPLATE is None:
    print(f"Warning: Excel template not found in any of the expected locations")
    EXCEL_TEMPLATE = Path("Loopbaan onderzoek 5.0.xlsx")  # Fallback


# ============================================================
# EXCEL FUNCTIONS
# ============================================================

def create_client_excel(client_id: str, client_data: Dict[str, Any]) -> bool:
    """
    Create a copy of the Excel template for a client and fill with client data.
    Returns True if successful, False otherwise.
    """
    try:
        if not EXCEL_TEMPLATE.exists():
            print(f"Excel template not found: {EXCEL_TEMPLATE}")
            return False

        excel_path = get_client_excel_path(client_id)

        # Copy template
        shutil.copy2(EXCEL_TEMPLATE, excel_path)

        # Update client info in the Excel file
        wb = load_workbook(excel_path)

        # Update client info in the Excel file
        fill_client_info(wb, client_data)

        # If there are questionnaire results, fill them too
        if client_data.get('big_five', {}).get('completed', False):
            fill_big_five_results(wb, client_data['big_five'])

        if client_data.get('career_anchors', {}).get('completed', False):
            fill_career_anchors_results(wb, client_data['career_anchors'])

        if client_data.get('career_clusters', {}).get('completed', False):
            fill_career_clusters_results(wb, client_data['career_clusters'])

        if client_data.get('culture', {}).get('completed', False):
            fill_culture_results(wb, client_data['culture'])

        if client_data.get('jcm', {}).get('completed', False):
            fill_jcm_results(wb, client_data['jcm'])

        wb.save(excel_path)
        return True
    except Exception as e:
        print(f"Error creating client Excel: {e}")
        return False


def fill_client_info(wb, client_data: Dict[str, Any]):
    """Fill client information in the Gegevens sheet"""
    if "Gegevens" not in wb.sheetnames:
        return

    ws = wb["Gegevens"]

    field_mapping = {
        'Cliënt': client_data.get('name', ''),
        'Email': client_data.get('email', ''),
        'Telefoon': client_data.get('telefoonnummer', ''),
        'Geboortedatum': client_data.get('geboortedatum', ''),
        'Land': client_data.get('land_van_herkomst', ''),
        'Leef situatie': client_data.get('leef_situatie', ''),
        'Opleidingen': client_data.get('opleidingen', ''),
        'Adres': client_data.get('adres', ''),
        'Toelichting': client_data.get('toelichting', ''),
    }

    for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=3):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                for label, value in field_mapping.items():
                    if label in cell.value:
                        ws.cell(row=cell.row, column=cell.column + 1, value=value)
                        break


def fill_big_five_results(wb, big_five_data: Dict[str, Any]):
    """Fill Big Five results in the Excel file"""
    sheet_name = None
    for name in wb.sheetnames:
        if "Big Five" in name or "1.1" in name:
            sheet_name = name
            break

    if not sheet_name:
        return

    ws = wb[sheet_name]

    dimensions = {
        'extraversie': 'Extraversie',
        'altruïsme': 'Altruïsme',
        'conciëntieusheid': 'Conciëntieusheid',
        'neuroticisme': 'Neuroticisme',
        'openheid': 'Openheid'
    }

    for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=10):
        for cell in row:
            if cell.value and "Score" in str(cell.value):
                score_row = cell.row + 1
                for col_idx in range(2, 7):
                    header = ws.cell(row=cell.row - 1, column=col_idx).value
                    if header:
                        for dim_key, dim_label in dimensions.items():
                            if dim_label in str(header):
                                ws.cell(row=score_row, column=col_idx,
                                        value=big_five_data.get(dim_key, {}).get('score', 0))
                break


def fill_career_anchors_results(wb, anchors_data: Dict[str, Any]):
    """Fill Career Anchors results in the Excel file"""
    sheet_name = None
    for name in wb.sheetnames:
        if "Loopbaanankers" in name or "2.0" in name:
            sheet_name = name
            break

    if not sheet_name:
        return

    ws = wb[sheet_name]

    anchor_mapping = {
        'V': 'Omhoog',
        'W': 'Veilig',
        'X': 'Vrij',
        'Y': 'Balans',
        'Z': 'Uitdaging'
    }

    for row in ws.iter_rows(min_row=1, max_row=100, min_col=1, max_col=10):
        for cell in row:
            if cell.value and "Totaal score" in str(cell.value):
                total_row = cell.row + 1
                for col_idx, (anchor, label) in enumerate(anchor_mapping.items(), start=2):
                    header = ws.cell(row=cell.row - 1, column=col_idx).value
                    if header and label in str(header):
                        ws.cell(row=total_row, column=col_idx, value=anchors_data.get('scores', {}).get(anchor, 0))
                break


def fill_career_clusters_results(wb, clusters_data: Dict[str, Any]):
    """Fill Career Clusters results in the Excel file"""
    sheet_name = None
    for name in wb.sheetnames:
        if "Carrière Clusters" in name or "Clusters" in name:
            sheet_name = name
            break

    if not sheet_name:
        return

    ws = wb[sheet_name]

    for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=10):
        for cell in row:
            if cell.value and "Totaal score" in str(cell.value):
                total_row = cell.row + 1
                for col_idx in range(2, 10):
                    cluster_id = str(col_idx - 1)
                    if cluster_id in clusters_data.get('scores', {}):
                        ws.cell(row=total_row, column=col_idx, value=clusters_data['scores'][cluster_id])
                break


def fill_culture_results(wb, culture_data: Dict[str, Any]):
    """Fill Culture Analysis results in the Excel file"""
    sheet_name = None
    for name in wb.sheetnames:
        if "Cultuur" in name and "analyse" in name.lower():
            sheet_name = name
            break

    if not sheet_name:
        return

    ws = wb[sheet_name]

    for row in ws.iter_rows(min_row=1, max_row=50, min_col=1, max_col=10):
        for cell in row:
            if cell.value and "Totaal score" in str(cell.value):
                culture_id = str(cell.row // 5)
                if culture_id in culture_data.get('scores', {}):
                    ws.cell(row=cell.row, column=3, value=culture_data['scores'][culture_id])
                break


def fill_jcm_results(wb, jcm_data: Dict[str, Any]):
    """Fill JCM results in the Excel file"""
    sheet_name = None
    for name in wb.sheetnames:
        if "J.C.M." in name or "JCM" in name:
            sheet_name = name
            break

    if not sheet_name:
        return

    ws = wb[sheet_name]

    component_mapping = {
        '1': 'Taakvaardigheid',
        '2': 'Taakidentiteit',
        '3': 'Taakbetekenis',
        '4': 'Autonomie',
        '5': 'Feedback'
    }

    for row in ws.iter_rows(min_row=1, max_row=30, min_col=1, max_col=4):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                for comp_id, comp_name in component_mapping.items():
                    if comp_name in cell.value:
                        if comp_id in jcm_data.get('answers', {}):
                            answer = jcm_data['answers'][comp_id]
                            if cell.column + 1 <= ws.max_column:
                                ws.cell(row=cell.row, column=cell.column + 1, value=answer)
                        break


def update_client_excel(client_id: str) -> bool:
    """
    Update an existing client's Excel file with latest data.
    Returns True if successful, False otherwise.
    """
    try:
        client = get_client(client_id)
        if not client:
            return False

        excel_path = get_client_excel_path(client_id)
        if not excel_path.exists():
            return create_client_excel(client_id, client)

        wb = load_workbook(excel_path)

        fill_client_info(wb, client)

        if client.get('big_five', {}).get('completed', False):
            fill_big_five_results(wb, client['big_five'])

        if client.get('career_anchors', {}).get('completed', False):
            fill_career_anchors_results(wb, client['career_anchors'])

        if client.get('career_clusters', {}).get('completed', False):
            fill_career_clusters_results(wb, client['career_clusters'])

        if client.get('culture', {}).get('completed', False):
            fill_culture_results(wb, client['culture'])

        if client.get('jcm', {}).get('completed', False):
            fill_jcm_results(wb, client['jcm'])

        wb.save(excel_path)
        return True
    except Exception as e:
        print(f"Error updating client Excel: {e}")
        return False


# ============================================================
# WORD REPORT FUNCTIONS
# ============================================================

def generate_word_report(client_id: str) -> bool:
    """
    Generate a Word document from the client's data.
    Returns True if successful, False otherwise.
    """
    try:
        client = get_client(client_id)
        if not client:
            return False

        word_path = get_client_word_path(client_id)

        doc = Document()

        setup_document_styles(doc)
        add_title(doc, client)
        add_client_info(doc, client)
        add_questionnaire_results(doc, client)
        add_footer(doc, client)

        doc.save(word_path)
        return True
    except Exception as e:
        print(f"Error generating Word report: {e}")
        return False


def setup_document_styles(doc):
    """Set up document styles"""
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)


def add_title(doc, client):
    """Add the report title"""
    title = doc.add_heading('Loopbaan Rapport', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(f'Voor: {client.get("name", "Onbekend")}')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.italic = True

    doc.add_paragraph()
    doc.add_paragraph('_' * 60)
    doc.add_paragraph()


def add_client_info(doc, client):
    """Add client information section"""
    doc.add_heading('Client Informatie', level=1)

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    client_fields = [
        ('Naam:', client.get('name', '')),
        ('Email:', client.get('email', '')),
        ('Telefoon:', client.get('telefoonnummer', '')),
        ('Geboortedatum:', client.get('geboortedatum', '')),
        ('Land van herkomst:', client.get('land_van_herkomst', '')),
        ('Leef situatie:', client.get('leef_situatie', '')),
        ('Opleidingen:', client.get('opleidingen', '')),
        ('Adres:', client.get('adres', '')),
    ]

    for label, value in client_fields:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[0].paragraphs[0].runs[0].bold = True
        row_cells[1].text = value if value else 'Niet ingevuld'

    doc.add_paragraph()


def add_questionnaire_results(doc, client):
    """Add all questionnaire results"""
    doc.add_heading('Vragenlijst Resultaten', level=1)

    if client.get('big_five', {}).get('completed', False):
        add_big_five_results(doc, client['big_five'])

    if client.get('career_anchors', {}).get('completed', False):
        add_career_anchors_results(doc, client['career_anchors'])

    if client.get('career_clusters', {}).get('completed', False):
        add_career_clusters_results(doc, client['career_clusters'])

    if client.get('culture', {}).get('completed', False):
        add_culture_results(doc, client['culture'])

    if client.get('jcm', {}).get('completed', False):
        add_jcm_results(doc, client['jcm'])


def add_big_five_results(doc, big_five):
    """Add Big Five results section"""
    doc.add_heading('Big Five Persoonlijkheidsdimensies', level=2)

    dimension_labels = {
        'extraversie': 'Extraversie',
        'altruïsme': 'Altruïsme',
        'conciëntieusheid': 'Conciëntieusheid',
        'neuroticisme': 'Neuroticisme',
        'openheid': 'Openheid'
    }

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'

    header_cells = table.rows[0].cells
    header_cells[0].text = 'Dimensie'
    header_cells[0].paragraphs[0].runs[0].bold = True
    header_cells[1].text = 'Score'
    header_cells[1].paragraphs[0].runs[0].bold = True
    header_cells[2].text = 'Niveau'
    header_cells[2].paragraphs[0].runs[0].bold = True

    for dim, label in dimension_labels.items():
        if dim in big_five:
            score = big_five[dim].get('score', 0)
            label_text = big_five[dim].get('label', 'Onbekend')

            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = f'{score}/50'
            row_cells[2].text = label_text

    doc.add_paragraph()


def add_career_anchors_results(doc, anchors):
    """Add Career Anchors results section"""
    doc.add_heading('Loopbaanankers', level=2)

    anchor_names = {
        'V': 'Omhoog komen',
        'W': 'Veilig voelen',
        'X': 'Vrij zijn',
        'Y': 'Balans vinden',
        'Z': 'Uitdaging zoeken'
    }

    if 'top1_name' in anchors:
        p = doc.add_paragraph()
        p.add_run('1e Anker: ').bold = True
        p.add_run(f"{anchors.get('top1_name', '')} ({anchors.get('top1_score', 0)} punten)")

    if 'top2_name' in anchors:
        p = doc.add_paragraph()
        p.add_run('2e Anker: ').bold = True
        p.add_run(f"{anchors.get('top2_name', '')} ({anchors.get('top2_score', 0)} punten)")

    if 'scores' in anchors:
        doc.add_paragraph('Alle ankers:')
        for anchor, score in anchors['scores'].items():
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{anchor_names.get(anchor, anchor)}: {score}/30")

    doc.add_paragraph()


def add_career_clusters_results(doc, clusters):
    """Add Career Clusters results section"""
    doc.add_heading('Carrière Clusters', level=2)

    if 'top1_name' in clusters:
        p = doc.add_paragraph()
        p.add_run('1e Cluster: ').bold = True
        p.add_run(f"{clusters.get('top1_name', '')} ({clusters.get('top1_score', 0)} matches)")

    if 'top2_name' in clusters:
        p = doc.add_paragraph()
        p.add_run('2e Cluster: ').bold = True
        p.add_run(f"{clusters.get('top2_name', '')} ({clusters.get('top2_score', 0)} matches)")

    if 'scores' in clusters and clusters['scores']:
        doc.add_paragraph('Alle clusters:')
        for cluster_id, score in sorted(clusters['scores'].items(), key=lambda x: x[1], reverse=True):
            from app.main import CAREER_CLUSTERS_DATA
            name = next((c.get('name', f'Cluster {c.get("id")}') for c in CAREER_CLUSTERS_DATA if
                         str(c.get('id')) == cluster_id), f'Cluster {cluster_id}')
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{name}: {score} matches")

    doc.add_paragraph()


def add_culture_results(doc, culture):
    """Add Culture Analysis results section"""
    doc.add_heading('Cultuur Analyse', level=2)

    culture_names = {
        '1': 'Mensgerichte cultuur',
        '2': 'Innovatieve cultuur',
        '3': 'Beheersgerichte cultuur',
        '4': 'Resultaatgerichte cultuur'
    }

    if 'top1_name' in culture:
        p = doc.add_paragraph()
        p.add_run('1e Cultuur: ').bold = True
        p.add_run(f"{culture.get('top1_name', '')} ({culture.get('top1_score', 0)} punten)")

    if 'top2_name' in culture:
        p = doc.add_paragraph()
        p.add_run('2e Cultuur: ').bold = True
        p.add_run(f"{culture.get('top2_name', '')} ({culture.get('top2_score', 0)} punten)")

    if 'scores' in culture:
        doc.add_paragraph('Alle culturen:')
        for cid, score in sorted(culture['scores'].items(), key=lambda x: x[1], reverse=True):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{culture_names.get(cid, f'Cultuur {cid}')}: {score}/20")

    doc.add_paragraph()


def add_jcm_results(doc, jcm):
    """Add JCM results section"""
    doc.add_heading('Werk Karakteristieken (JCM)', level=2)

    component_names = {
        '1': 'Taakvaardigheid',
        '2': 'Taakidentiteit',
        '3': 'Taakbetekenis',
        '4': 'Autonomie',
        '5': 'Feedback'
    }

    component_letters = {
        '1': 'a',
        '2': 'b',
        '3': 'c',
        '4': 'd',
        '5': 'e'
    }

    if 'answers' in jcm:
        for comp_id, answer in jcm['answers'].items():
            name = component_names.get(comp_id, f'Component {comp_id}')
            letter = component_letters.get(comp_id, '')

            doc.add_heading(f'{name} ({letter})', level=3)
            p = doc.add_paragraph(answer if answer else 'Geen antwoord ingevuld')
            p.style = 'Normal'
            doc.add_paragraph()


def add_footer(doc, client):
    """Add footer with timestamp and generation info"""
    doc.add_paragraph('_' * 60)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p.add_run(f'Rapport gegenereerd op: {datetime.now().strftime("%d-%m-%Y %H:%M")}')
    p.add_run('\n')
    p.add_run(f'Klant: {client.get("name", "Onbekend")} (ID: {client.get("id", "")})')
    p.add_run('\n')
    p.add_run('Loopbaan Onderzoek - Automatisch gegenereerd rapport')

    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(128, 128, 128)